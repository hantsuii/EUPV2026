import importlib.util
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("po_check", ROOT / "scripts" / "po_check.py")
po_check = importlib.util.module_from_spec(SPEC)
assert SPEC.loader
SPEC.loader.exec_module(po_check)


def save_book(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)


class PoCheckTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.purchase = self.base / "purchase.xlsx"
        self.odp = self.base / "odp.xlsx"
        self.warehouse = self.base / "warehouse.xlsx"
        self.output = self.base / "result.xlsx"
        self.json_output = self.base / "result.json"

        head_headers = [
            "Purchase order Number", "Purchase Type", "TMS bill Number / voucher bill Number", "Customer Po",
            "Category", "Estimated time of arrival", "Status", "Approval Status", "Note", "X",
        ]
        head_rows = [
            ["PO1", "Internal", "T1", "ABC123.1", "PV", date(2026, 1, 10), "In Stock", "", "Approved", ""],
            ["PO2", "Offshore purchase", "T2", "REF456.2", "PV", date(2026, 2, 1), "Open", "", "Approving", ""],
            ["PO3", "Internal", "T3", "", "PV", date(2026, 3, 1), "Open", "", "Approved", ""],
            ["PO4", "Internal", "T4", "MUL789.1", "PV", date(2026, 4, 1), "Open", "", "Approved", ""],
        ]
        line_headers = ["Purchase order Number", "Customer Model", "Qty", "Storage Space"]
        line_rows = [
            ["PO1", "SKU1", 2, "WH1"], ["PO1", "SKU1", 3, "WH1"],
            ["PO2", "SKU2", 4, "WH1"], ["PO3", "SKU3", 1, "WH1"], ["PO4", "SKU4", 1, "WH1"],
        ]
        odp_headers = [
            "TCL REFERENCE", "New Ark PO Type", "New Ark WH", "New Ark SKU", "Model", "QUANTITY",
            "SPTN-PVHK New Ark PO#", "NewArk ETA PO", "ETA for New Ark update",
        ]
        odp_rows = [
            ["ABC123.1", "Internal", "Warehouse One", "SKU1", "Model 1", 5, "PO1", date(2026, 1, 11), date(2026, 1, 10)],
            ["REF456.2", "Offshore", "Warehouse One", "SKU2", "Model 2", 4, "PO2", "", date(2026, 2, 5)],
            ["MUL789.1", "Internal", "Warehouse One", "SKU4", "Model 4", 1, "PO4", date(2026, 4, 1), date(2026, 4, 1)],
            ["MUL789.1", "Internal", "Warehouse One", "SKU4", "Model 4", 0, "PO4", date(2026, 4, 1), date(2026, 4, 3)],
        ]
        save_book(self.purchase, {"purchase order head": [head_headers, *head_rows], "purchase order line": [line_headers, *line_rows]})
        save_book(self.odp, {"PO Check": [odp_headers, *odp_rows]})
        save_book(self.warehouse, {"Sheet1": [["Virtual warehouse name", "Virtual warehouse code"], ["Warehouse One", "WH1"]]})

    def tearDown(self):
        self.temp.cleanup()

    def test_business_rules_and_priority(self):
        report = po_check.build_report(self.purchase, self.odp, self.warehouse)
        by_po = {row["po"]: row for row in report["results"]}
        self.assertEqual(by_po["PO1"]["overall"], "通过")
        self.assertEqual(by_po["PO1"]["etaPoResult"], "通过")
        self.assertEqual(by_po["PO1"]["skuResult"], "通过")
        self.assertEqual(by_po["PO1"]["warehouseResult"], "通过")
        self.assertEqual(by_po["PO2"]["etaPoResult"], "不适用")
        self.assertTrue(by_po["PO2"]["etaNeedsAdjustment"])
        self.assertEqual(by_po["PO3"]["overall"], "需复核")
        self.assertEqual(by_po["PO4"]["etaUpdateResult"], "需复核：多条 ETA 更新")
        self.assertEqual([row["po"] for row in report["results"]], ["PO3", "PO4", "PO2", "PO1"])
        self.assertEqual(report["summary"], {"checked": 4, "approved": 3, "approving": 1, "otherApprovalStatus": 0, "issuesOrReview": 3, "etaAdjustments": 1})

    def test_output_preserves_eta_and_in_stock(self):
        report = po_check.build_report(self.purchase, self.odp, self.warehouse)
        po_check.write_output(report, self.output, "zh")
        workbook = load_workbook(self.output, data_only=True)
        self.assertEqual(workbook.sheetnames, ["PO Details", "PO Check Results", "SKU Qty Results", "Summary"])
        details = workbook["PO Details"]
        self.assertEqual(details.max_row, 5)
        self.assertEqual(details["A2"].value, "PO1")
        self.assertEqual(details["F2"].value.date(), date(2026, 1, 10))
        self.assertEqual(details["G2"].value, "In Stock")
        self.assertEqual(details["H2"].value, "Approved")

    def test_cli_writes_json(self):
        code = po_check.main([
            "--purchase", str(self.purchase), "--odp", str(self.odp), "--warehouse", str(self.warehouse),
            "--output", str(self.output), "--json-output", str(self.json_output),
        ])
        self.assertEqual(code, 0)
        payload = json.loads(self.json_output.read_text(encoding="utf-8"))
        self.assertEqual(payload["summary"]["checked"], 4)


if __name__ == "__main__":
    unittest.main()
