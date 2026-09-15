#!/usr/bin/env python3
"""Portable reproduction of the EUPV2026 PO Check calculation engine."""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: install with 'pip install -r requirements.txt'.") from exc


TYPE_MAP = {
    "INTERNAL": "Internal",
    "OFFSHORE": "Offshore purchase",
    "ONGOING B/L CHANGE": "Internal",
}
APPROVAL_VALUES = {"APPROVED", "APPROVING"}
RESULT_TRANSLATIONS = {
    "通过": "Passed", "需复核": "Review required", "有差异": "Difference", "不适用": "N/A",
    "需定义映射": "Mapping required", "未在 PO Check 找到": "Not found in PO Check", "需人工核对": "Manual review",
    "NewArk ETA PO 不一致": "NewArk ETA PO mismatch", "需复核：多条 NewArk ETA PO": "Review: multiple NewArk ETA PO values",
    "NewArk ETA PO 缺失": "NewArk ETA PO missing", "ETA 更新不一致": "Updated ETA mismatch",
    "需复核：多条 ETA 更新": "Review: multiple updated ETA values", "ETA 更新缺失": "Updated ETA missing",
    "建议更新 ETA": "Update ETA", "需人工选择 ETA": "Select ETA manually", "无需更新 ETA": "No ETA update",
    "采购类型不一致": "Purchase type mismatch", "需复核：多种 PO 类型": "Review: multiple PO types",
    "TCL Reference 不一致": "TCL Reference mismatch", "采购单缺少 SKU": "SKU missing from purchase order",
    "PO Check 缺少 SKU": "SKU missing from PO Check", "数量不一致": "Quantity mismatch",
    "仓库映射缺失": "Warehouse mapping missing", "仓库不一致": "Warehouse mismatch", "（空）": "(blank)",
}
PO_DETAILS_HEADERS = [
    "Purchase order Number", "Purchase Type", "TMS bill Number / voucher bill Number", "Customer Po",
    "Category", "Estimated time of arrival", "Status", "Approval Status", "Note",
]
PO_RESULT_HEADERS = [
    "PO", "Approval status", "Overall result", "Purchase Status", "SKU", "Model", "Purchase Type",
    "Expected Purchase Type", "ODP PO Type", "Purchase ETA", "NewArk ETA PO", "NewArk ETA PO check",
    "ETA for New Ark update", "Update ETA check", "Customer Po TCL Ref", "X column TCL Ref",
    "PO Check TCL Ref", "TCL Ref check", "SKU / Qty check", "SKU / Qty detail", "Purchase warehouse",
    "PO Check New Ark WH", "Warehouse check", "Suggested ETA", "ETA action", "Notes",
]


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", text(value)).upper()


def unique(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def cell(row: list[Any] | tuple[Any, ...], index: int) -> Any:
    return row[index] if 0 <= index < len(row) else ""


def header_index(headers: list[Any], name: str) -> int:
    wanted = normalize(name)
    return next((i for i, value in enumerate(headers) if normalize(value) == wanted), -1)


def workbook_rows(path: Path, sheet_name: str) -> list[list[Any]]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"Unsupported workbook format: {path.suffix}; use .xlsx or .xlsm")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet_name}")
        return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()


def date_iso(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (date(1899, 12, 30) + timedelta(days=float(value))).isoformat()
        except (OverflowError, ValueError):
            return text(value)
    value_text = text(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value_text):
        return value_text
    match = re.search(r"(\d{4})[/.](\d{1,2})[/.](\d{1,2})", value_text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    for pattern in ("%m/%d/%Y", "%d/%m/%Y", "%Y%m%d", "%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(value_text, pattern).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(value_text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return value_text


def parsed_date(value: Any) -> date | None:
    iso = date_iso(value)
    try:
        return date.fromisoformat(iso)
    except ValueError:
        return None


def eta_within_one_day(left: Any, right: Any) -> bool:
    left_date, right_date = parsed_date(left), parsed_date(right)
    return bool(left and right and left_date and right_date and abs((left_date - right_date).days) <= 1)


def eta_values_within_one_day(values: Iterable[Any]) -> bool:
    nonblank = [value for value in values if text(value)]
    dates = [parsed_date(value) for value in nonblank]
    if any(value is None for value in dates):
        return False
    if len(dates) < 2:
        return True
    valid_dates = [value for value in dates if value is not None]
    return (max(valid_dates) - min(valid_dates)).days <= 1


def number(value: Any) -> float:
    try:
        result = float(text(value).replace(",", ""))
        return result if math.isfinite(result) else 0.0
    except ValueError:
        return 0.0


def values_label(values: list[str]) -> str:
    return " | ".join(values) if values else "（空）"


def choose_approval_column(rows: list[list[Any]], headers: list[Any]) -> int:
    max_columns = max((len(row) for row in rows[:100]), default=0)
    best, best_count = -1, -1
    for column in range(max_columns):
        count = sum(normalize(cell(row, column)) in APPROVAL_VALUES for row in rows)
        if count > best_count:
            best, best_count = column, count
    if best_count < 1:
        note_column = header_index(headers, "Note")
        if note_column >= 0:
            return note_column
        approval_header = header_index(headers, "Approval Status")
        if approval_header >= 0:
            return approval_header + 1
        raise ValueError("Could not identify the Approved / Approving column")
    return best


def references_from(value: Any, known_refs: list[str]) -> list[str]:
    source = re.sub(r"[，、；;|/\\\n\r]+", " ", normalize(value))
    hits = [reference for reference in known_refs if normalize(reference) in source]
    generic = re.findall(r"(?<![A-Z0-9])(?:\d{2}[A-Z]{3}\d{3}(?:\.\d{1,3})+|[A-Z]{3,8}\d{3}(?:\.\d{1,3})+)(?![A-Z0-9])", source)
    return sorted(unique([*hits, *generic]))


def required_indexes(headers: list[Any], names: list[str], group: str) -> dict[str, int]:
    indexes = {name: header_index(headers, name) for name in names}
    missing = [name for name, index in indexes.items() if index < 0]
    if missing:
        raise ValueError(f"{group} missing required columns: {', '.join(missing)}")
    return indexes


def build_warehouse_map(rows: list[list[Any]]) -> dict[str, str]:
    if len(rows) < 2:
        raise ValueError("Warehouse mapping contains no data")
    headers = rows[0]
    indexes = required_indexes(headers, ["Virtual warehouse name", "Virtual warehouse code"], "Warehouse mapping")
    result: dict[str, str] = {}
    for row in rows[1:]:
        name = text(cell(row, indexes["Virtual warehouse name"]))
        code = text(cell(row, indexes["Virtual warehouse code"]))
        if name:
            result[normalize(name)] = name
        if code and name:
            result[normalize(code)] = name
    return result


def result_priority(row: dict[str, Any]) -> int:
    if row["overall"] == "需复核":
        return 0
    if row["etaNeedsAdjustment"]:
        return 1
    if row["overall"] == "有差异":
        return 2
    return 3


def sorted_results(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(rows, key=lambda row: (result_priority(row), row["po"]))


def build_report(purchase_path: Path, odp_path: Path, warehouse_path: Path) -> dict[str, Any]:
    head_rows = workbook_rows(purchase_path, "purchase order head")
    line_rows = workbook_rows(purchase_path, "purchase order line")
    odp_rows = workbook_rows(odp_path, "PO Check")
    warehouse_rows = workbook_rows(warehouse_path, "Sheet1")
    if len(head_rows) < 2 or len(line_rows) < 2 or len(odp_rows) < 2:
        raise ValueError("Input workbooks contain no checkable data")

    head_headers, line_headers, odp_headers = head_rows[0], line_rows[0], odp_rows[0]
    head_data = [row for row in head_rows[1:] if text(cell(row, 1))]
    line_data = [row for row in line_rows[1:] if text(cell(row, 0))]
    odp_data = [row for row in odp_rows[1:] if text(cell(row, 0))]
    head_names = ["Purchase order Number", "Purchase Type", "TMS bill Number / voucher bill Number", "Customer Po", "Category", "Estimated time of arrival", "Status"]
    line_names = ["Purchase order Number", "Customer Model", "Qty", "Storage Space"]
    odp_names = ["TCL REFERENCE", "New Ark PO Type", "New Ark WH", "New Ark SKU", "Model", "QUANTITY", "SPTN-PVHK New Ark PO#", "NewArk ETA PO", "ETA for New Ark update"]
    hi = required_indexes(head_headers, head_names, "purchase order head")
    li = required_indexes(line_headers, line_names, "purchase order line")
    oi = required_indexes(odp_headers, odp_names, "PO Check")
    approval_column = choose_approval_column(head_data, head_headers)
    x_column = approval_column + 1
    warehouse_map = build_warehouse_map(warehouse_rows)

    odp_by_po: dict[str, list[list[Any]]] = defaultdict(list)
    for row in odp_data:
        po = text(cell(row, oi["SPTN-PVHK New Ark PO#"]))
        if po:
            odp_by_po[po].append(row)
    lines_by_po: dict[str, list[list[Any]]] = defaultdict(list)
    for row in line_data:
        po = text(cell(row, li["Purchase order Number"]))
        if po:
            lines_by_po[po].append(row)

    results: list[dict[str, Any]] = []
    sku_results: list[dict[str, Any]] = []
    po_output_rows: list[list[Any]] = []

    for head_row in head_data:
        po = text(cell(head_row, hi["Purchase order Number"]))
        approval = text(cell(head_row, approval_column))
        purchase_status = text(cell(head_row, hi["Status"]))
        purchase_eta = date_iso(cell(head_row, hi["Estimated time of arrival"]))
        checks = odp_by_po.get(po, [])
        odp_types = unique(text(cell(row, oi["New Ark PO Type"])) for row in checks)
        expected_types = unique(TYPE_MAP.get(normalize(value), "需定义映射") for value in odp_types)
        eta_po = unique(date_iso(cell(row, oi["NewArk ETA PO"])) for row in checks)
        eta_update = unique(date_iso(cell(row, oi["ETA for New Ark update"])) for row in checks)
        odp_refs = sorted(unique(normalize(cell(row, oi["TCL REFERENCE"])) for row in checks))
        skus = sorted(unique(text(cell(row, oi["New Ark SKU"])) for row in checks))
        models = sorted(unique(text(cell(row, oi["Model"])) for row in checks))
        customer_refs = references_from(cell(head_row, hi["Customer Po"]), odp_refs)
        x_refs = references_from(cell(head_row, x_column), odp_refs)
        source_refs = sorted(unique([*customer_refs, *x_refs]))
        type_result = eta_po_result = eta_update_result = ref_result = sku_result = warehouse_result = "不适用"
        suggested_eta, action = "", "不适用"
        notes: list[str] = []

        purchase_warehouses = sorted(unique(
            warehouse_map.get(normalize(cell(line, li["Storage Space"])))
            or (f"未映射：{text(cell(line, li['Storage Space']))}" if text(cell(line, li["Storage Space"])) else "")
            for line in lines_by_po.get(po, [])
        ))
        odp_warehouses = sorted(unique(text(cell(row, oi["New Ark WH"])) for row in checks))

        if not checks:
            type_result = eta_po_result = eta_update_result = ref_result = sku_result = warehouse_result = "未在 PO Check 找到"
            action = "需人工核对"
            notes.append("采购单号未在 PO Check 找到。")
        else:
            type_result = ("通过" if text(cell(head_row, hi["Purchase Type"])) == expected_types[0] else "采购类型不一致") if len(expected_types) == 1 else "需复核：多种 PO 类型"
            if normalize(approval) == "APPROVING":
                eta_po_result = "不适用"
            elif len(eta_po) == 1:
                eta_po_result = "通过" if eta_within_one_day(eta_po[0], purchase_eta) else "NewArk ETA PO 不一致"
            else:
                eta_po_result = "需复核：多条 NewArk ETA PO" if eta_po else "NewArk ETA PO 缺失"
            if len(eta_update) == 1:
                eta_update_result = "通过" if eta_values_within_one_day([purchase_eta, *eta_po, eta_update[0]]) else "ETA 更新不一致"
            else:
                eta_update_result = "需复核：多条 ETA 更新" if eta_update else "ETA 更新缺失"
            ref_result = "通过" if source_refs == odp_refs else "TCL Reference 不一致"
            purchase_sku: dict[str, float] = defaultdict(float)
            odp_sku: dict[str, float] = defaultdict(float)
            for line in lines_by_po.get(po, []):
                sku = text(cell(line, li["Customer Model"]))
                if sku:
                    purchase_sku[sku] += number(cell(line, li["Qty"]))
            for odp_row in checks:
                sku = text(cell(odp_row, oi["New Ark SKU"]))
                if sku:
                    odp_sku[sku] += number(cell(odp_row, oi["QUANTITY"]))
            sku_issue_count = 0
            for sku in sorted(set(purchase_sku) | set(odp_sku)):
                left, right = purchase_sku.get(sku, 0.0), odp_sku.get(sku, 0.0)
                sku_status = "采购单缺少 SKU" if sku not in purchase_sku else "PO Check 缺少 SKU" if sku not in odp_sku else "通过" if abs(left - right) < 0.000001 else "数量不一致"
                sku_issue_count += sku_status != "通过"
                sku_results.append({"po": po, "sku": sku, "purchaseQty": left, "odpQty": right, "result": sku_status})
            sku_result = f"SKU/数量有 {sku_issue_count} 项差异" if sku_issue_count else "通过"
            if purchase_warehouses and not any(value.startswith("未映射：") for value in purchase_warehouses) and odp_warehouses and sorted(map(normalize, purchase_warehouses)) == sorted(map(normalize, odp_warehouses)):
                warehouse_result = "通过"
            elif any(value.startswith("未映射：") for value in purchase_warehouses):
                warehouse_result = "仓库映射缺失"
            else:
                warehouse_result = "仓库不一致"
            if len(eta_update) == 1 and eta_update_result == "ETA 更新不一致":
                suggested_eta, action = eta_update[0], "建议更新 ETA"
            elif len(eta_update) > 1:
                action = "需人工选择 ETA"
            else:
                action = "无需更新 ETA" if eta_update_result == "通过" else "需人工核对"

        check_results = [type_result, eta_po_result, eta_update_result, ref_result, sku_result, warehouse_result]
        overall = "通过" if all(value in {"通过", "不适用"} for value in check_results) else "需复核" if any("需复核" in value or "未在" in value for value in check_results) else "有差异"
        eta_needs_adjustment = eta_update_result == "ETA 更新不一致" and overall == "有差异"
        this_po_skus = [item for item in sku_results if item["po"] == po]
        sku_detail = "\n".join(f"{item['sku']}: {format_qty(item['purchaseQty'])} / {format_qty(item['odpQty'])} ({item['result']})" for item in this_po_skus) or "（空）"
        results.append({
            "po": po, "approval": approval, "purchaseStatus": purchase_status, "skus": values_label(skus), "models": values_label(models),
            "purchaseType": text(cell(head_row, hi["Purchase Type"])), "expectedType": values_label(expected_types), "odpTypes": values_label(odp_types),
            "purchaseEta": purchase_eta, "etaPo": values_label(eta_po), "etaUpdate": values_label(eta_update), "typeResult": type_result,
            "etaPoResult": eta_po_result, "etaUpdateResult": eta_update_result, "customerRefs": values_label(customer_refs), "xRefs": values_label(x_refs),
            "odpRefs": values_label(odp_refs), "refResult": ref_result, "skuResult": sku_result, "skuQtyDetail": sku_detail,
            "purchaseWarehouses": values_label(purchase_warehouses), "odpWarehouses": values_label(odp_warehouses), "warehouseResult": warehouse_result,
            "suggestedEta": suggested_eta, "action": action, "overall": overall, "etaNeedsAdjustment": eta_needs_adjustment, "notes": " ".join(notes),
        })
        po_output_rows.append([
            po, text(cell(head_row, hi["Purchase Type"])), text(cell(head_row, hi["TMS bill Number / voucher bill Number"])),
            text(cell(head_row, hi["Customer Po"])), text(cell(head_row, hi["Category"])), purchase_eta, purchase_status,
            approval, text(cell(head_row, x_column)),
        ])

    results = sorted_results(results)
    return {"results": results, "skuResults": sku_results, "poOutputRows": po_output_rows, "summary": summary_counts(results)}


def format_qty(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else str(value)


def summary_counts(results: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "checked": len(results),
        "approved": sum(normalize(row["approval"]) == "APPROVED" for row in results),
        "approving": sum(normalize(row["approval"]) == "APPROVING" for row in results),
        "otherApprovalStatus": sum(normalize(row["approval"]) not in APPROVAL_VALUES for row in results),
        "issuesOrReview": sum(row["overall"] != "通过" for row in results),
        "etaAdjustments": sum(bool(row["etaNeedsAdjustment"]) for row in results),
    }


def display(value: Any, language: str) -> Any:
    if language != "en" or not isinstance(value, str):
        return value
    match = re.fullmatch(r"SKU/数量有 (\d+) 项差异", value)
    return f"{match.group(1)} SKU / quantity difference(s)" if match else RESULT_TRANSLATIONS.get(value, value)


def write_output(report: dict[str, Any], output_path: Path, language: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EUPV2026 PO Check"
    header_fill = PatternFill("solid", fgColor="2F75B5")
    header_font = Font(bold=True, color="FFFFFF")
    row_fills = {"review": PatternFill("solid", fgColor="E4D9F6"), "eta": PatternFill("solid", fgColor="C6EFCE"), "issue": PatternFill("solid", fgColor="FFE9CC")}
    ordered = report["results"]
    po_rows = [PO_RESULT_HEADERS] + [[
        row["po"], row["approval"], display(row["overall"], language), row["purchaseStatus"], row["skus"], row["models"], row["purchaseType"],
        row["expectedType"], row["odpTypes"], row["purchaseEta"], row["etaPo"], display(row["etaPoResult"], language), row["etaUpdate"],
        display(row["etaUpdateResult"], language), row["customerRefs"], row["xRefs"], row["odpRefs"], display(row["refResult"], language),
        display(row["skuResult"], language), row["skuQtyDetail"], row["purchaseWarehouses"], row["odpWarehouses"], display(row["warehouseResult"], language),
        row["suggestedEta"], display(row["action"], language), row["notes"],
    ] for row in ordered]
    sku_rows = [["PO", "SKU", "Purchase Qty", "PO Check Qty", "Result"]] + [[r["po"], r["sku"], r["purchaseQty"], r["odpQty"], display(r["result"], language)] for r in report["skuResults"]]
    summary = report["summary"]
    summary_rows = [["Metric", "Value"], ["Checked POs", summary["checked"]], ["Approved", summary["approved"]], ["Approving", summary["approving"]],
                    ["Other approval status", summary["otherApprovalStatus"]], ["Issues / Review", summary["issuesOrReview"]], ["ETA adjustments", summary["etaAdjustments"]],
                    ["Check scope", "All PO statuses; all checks enabled"], ["PO Details ETA", "Original purchase-order ETA"],
                    ["Type mapping", "Internal → Internal"], ["Type mapping", "Offshore → Offshore purchase"], ["Type mapping", "ongoing B/L change → Internal"]]
    sheet_rows = [("PO Details", [PO_DETAILS_HEADERS, *report["poOutputRows"]]), ("PO Check Results", po_rows), ("SKU Qty Results", sku_rows), ("Summary", summary_rows)]
    for name, rows in sheet_rows:
        ws = workbook.create_sheet(name)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for item in ws[1]:
            item.fill, item.font = header_fill, header_font
            item.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for index in range(1, ws.max_column + 1):
            sample = [text(ws.cell(row=row, column=index).value) for row in range(1, min(ws.max_row, 100) + 1)]
            ws.column_dimensions[get_column_letter(index)].width = min(42, max(12, *(len(value) + 2 for value in sample)))
        for row in ws.iter_rows(min_row=2):
            for item in row:
                item.alignment = Alignment(vertical="top", wrap_text="\n" in text(item.value))
        if name == "PO Details":
            for row_index in range(2, ws.max_row + 1):
                iso = text(ws.cell(row_index, 6).value)
                parsed = parsed_date(iso)
                if parsed:
                    ws.cell(row_index, 6).value = parsed
                    ws.cell(row_index, 6).number_format = "yyyy-mm-dd"
                po = text(ws.cell(row_index, 1).value)
                result = next((item for item in ordered if item["po"] == po), None)
                if result and result["etaNeedsAdjustment"]:
                    ws.cell(row_index, 6).fill = row_fills["eta"]
        elif name == "PO Check Results":
            for row_index, result in enumerate(ordered, 2):
                fill = row_fills["review"] if result["overall"] == "需复核" else row_fills["eta"] if result["etaNeedsAdjustment"] else row_fills["issue"] if result["overall"] == "有差异" else None
                if fill:
                    for item in ws[row_index]:
                        item.fill = fill
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    skill_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--purchase", required=True, type=Path, help="Purchase-order .xlsx/.xlsm")
    parser.add_argument("--odp", required=True, type=Path, help="EUPV ODP MASTER .xlsx/.xlsm")
    parser.add_argument("--warehouse", type=Path, default=skill_root / "assets" / "logical_warehouse_list.xlsx")
    parser.add_argument("--output", required=True, type=Path, help="Output .xlsx")
    parser.add_argument("--json-output", type=Path, help="Optional machine-readable report")
    parser.add_argument("--language", choices=("zh", "en"), default="zh")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        for label, path in (("purchase", args.purchase), ("odp", args.odp), ("warehouse", args.warehouse)):
            if not path.is_file():
                raise FileNotFoundError(f"{label} workbook not found: {path}")
        report = build_report(args.purchase, args.odp, args.warehouse)
        write_output(report, args.output, args.language)
        if args.json_output:
            args.json_output.parent.mkdir(parents=True, exist_ok=True)
            args.json_output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        summary = report["summary"]
        print(json.dumps({"output": str(args.output), **summary}, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(f"PO Check failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
